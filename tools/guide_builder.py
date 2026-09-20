"""Оформление практического руководства в формате Word."""

from io import BytesIO
from pathlib import Path

from PIL import Image
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent.parent

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(13)
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x99)


class Guide:
    """Документ-руководство: заголовки, шаги, врезки и рисунки."""

    def __init__(self, title: str, subtitle: str = ""):
        self.document = docx.Document()
        self._setup()
        self._cover(title, subtitle)

    # ------------------------------------------------------------------ #

    def _setup(self) -> None:
        section = self.document.sections[0]
        section.page_width = Twips(11910)
        section.page_height = Twips(16840)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(1.6)

        style = self.document.styles["Normal"]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), BODY_FONT)
        fonts.set(qn("w:cs"), BODY_FONT)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15

    def _paragraph(self, align=None, left=None, space_before=None, space_after=None):
        paragraph = self.document.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        if left is not None:
            paragraph.paragraph_format.left_indent = left
        if space_before is not None:
            paragraph.paragraph_format.space_before = space_before
        if space_after is not None:
            paragraph.paragraph_format.space_after = space_after
        return paragraph

    def _cover(self, title: str, subtitle: str) -> None:
        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(2))
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(19)
        if subtitle:
            paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(14))
            run = paragraph.add_run(subtitle)
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ------------------------------------------------------------------ #
    # Блоки текста
    # ------------------------------------------------------------------ #

    def h1(self, text: str) -> None:
        paragraph = self._paragraph(space_before=Pt(16), space_after=Pt(6))
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = ACCENT

    def h2(self, text: str) -> None:
        paragraph = self._paragraph(space_before=Pt(11), space_after=Pt(4))
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(13.5)

    def text(self, *parts) -> None:
        """Абзац. Части — строки или пары (текст, стиль): 'b', 'i', 'c'."""
        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.JUSTIFY)
        self._fill(paragraph, parts)

    def _fill(self, paragraph, parts) -> None:
        for part in parts:
            if isinstance(part, tuple):
                content, style = part
                run = paragraph.add_run(content)
                if "b" in style:
                    run.bold = True
                if "i" in style:
                    run.italic = True
                if "c" in style:
                    run.font.name = CODE_FONT
                    run.font.size = Pt(11)
                    run.element.rPr.rFonts.set(qn("w:cs"), CODE_FONT)
            else:
                paragraph.add_run(part)

    def step(self, number: int, *parts) -> None:
        paragraph = self._paragraph(left=Cm(0.8), space_after=Pt(3))
        paragraph.paragraph_format.first_line_indent = Cm(-0.8)
        run = paragraph.add_run(f"{number}. ")
        run.bold = True
        self._fill(paragraph, parts)

    def bullet(self, *parts) -> None:
        paragraph = self._paragraph(left=Cm(0.8), space_after=Pt(3))
        paragraph.paragraph_format.first_line_indent = Cm(-0.4)
        paragraph.add_run("— ")
        self._fill(paragraph, parts)

    def note(self, label: str, *parts) -> None:
        """Врезка: выделенное замечание с подписью."""
        paragraph = self._paragraph(left=Cm(0.5), space_before=Pt(6), space_after=Pt(6))
        run = paragraph.add_run(f"{label} ")
        run.bold = True
        run.font.color.rgb = ACCENT
        self._fill(paragraph, parts)

    def qa(self, question: str, answer_parts) -> None:
        paragraph = self._paragraph(left=Cm(0.5), space_before=Pt(7), space_after=Pt(2))
        paragraph.paragraph_format.first_line_indent = Cm(-0.5)
        run = paragraph.add_run("? ")
        run.bold = True
        run.font.color.rgb = ACCENT
        run = paragraph.add_run(question)
        run.bold = True

        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.JUSTIFY, left=Cm(0.5), space_after=Pt(2))
        self._fill(paragraph, answer_parts)

    def empty(self, count: int = 1) -> None:
        for _ in range(count):
            self._paragraph(space_after=Pt(0))

    def page_break(self) -> None:
        self.document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ #
    # Рисунки
    # ------------------------------------------------------------------ #

    def picture(self, images, caption: str = "", width_cm: float = 16.0,
                max_height_cm: float = 17.0) -> None:
        composed = _compose([Path(image) for image in images])
        ratio = composed.height / composed.width

        width = Cm(width_cm)
        height = Cm(width_cm * ratio)
        if height.cm > max_height_cm:
            height = Cm(max_height_cm)
            width = Cm(max_height_cm / ratio)

        buffer = BytesIO()
        composed.save(buffer, format="PNG")
        buffer.seek(0)

        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, space_before=Pt(6),
                                    space_after=Pt(2))
        paragraph.add_run().add_picture(buffer, width=width, height=height)

        if caption:
            paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER, space_after=Pt(8))
            run = paragraph.add_run(caption)
            run.italic = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def save(self, *paths: Path) -> None:
        for path in paths:
            path.parent.mkdir(parents=True, exist_ok=True)
            self.document.save(str(path))
            print(f"Готово: {path}")


def _compose(images: list[Path]) -> Image.Image:
    opened = [Image.open(path).convert("RGB") for path in images]
    if len(opened) == 1:
        return opened[0]

    height = max(image.height for image in opened)
    scaled = [
        image if image.height == height
        else image.resize((round(image.width * height / image.height), height), Image.LANCZOS)
        for image in opened
    ]
    gap = round(height * 0.02)
    total = sum(image.width for image in scaled) + gap * (len(scaled) - 1)
    canvas = Image.new("RGB", (total, height), "white")

    offset = 0
    for image in scaled:
        canvas.paste(image, (offset, 0))
        offset += image.width + gap
    return canvas


def app(lab: str, *names: str) -> list[Path]:
    """Снимки экрана приложения."""
    return [ROOT / lab / "screens" / f"{name}.png" for name in names]


def site(*names: str) -> list[Path]:
    """Снимки страниц GitHub."""
    return [ROOT / "docs" / "guide" / f"{name}.png" for name in names]
