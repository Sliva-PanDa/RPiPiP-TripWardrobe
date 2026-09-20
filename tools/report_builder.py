"""Сборка отчётов по лабораторным работам в формате Word.

Оформление повторяет образец кафедры: Times New Roman 14 пт, выравнивание
по ширине, абзацный отступ 1,25 см, листинги 12 пт, подрисуночные подписи
по центру.
"""

import subprocess
from io import BytesIO
from pathlib import Path

from PIL import Image
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips

ROOT = Path(__file__).resolve().parent.parent

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(14)
CODE_SIZE = Pt(12)
FIRST_LINE_INDENT = Cm(1.25)

# Ширина изображения с парой снимков iPhone и с парой снимков iPad.
PHONE_PAIR_WIDTH = Cm(14.48)
PHONE_PAIR_MAX_HEIGHT = Cm(16.0)
TABLET_PAIR_WIDTH = Cm(16.5)

UNIVERSITY = "ГОМЕЛЬСКИЙ ГОСУДАРСТВЕННЫЙ ТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ ИМЕНИ П. О. СУХОГО"
FACULTY = ("Факультет автоматизированных и информационных систем "
           "Кафедра «Информационные технологии»")
DISCIPLINE = "Разработка приложений для iPhone и iPad"
STUDENT = "Выполнил: студент гр. ИТП-41 Сеноженский В. В."
TEACHER = "Шаблинский Д. А."
CITY_YEAR = "Гомель 2026"


class Report:
    """Документ отчёта по одной лабораторной работе."""

    def __init__(self, number: int, topic_lines: list[str]):
        self.document = docx.Document()
        self._setup_page()
        self._setup_styles()
        self._title_page(number, topic_lines)

    # ------------------------------------------------------------------ #
    # Базовое оформление
    # ------------------------------------------------------------------ #

    def _setup_page(self) -> None:
        section = self.document.sections[0]
        section.page_width = Twips(11910)
        section.page_height = Twips(16840)
        section.top_margin = Twips(1040)
        section.right_margin = Twips(708)
        section.bottom_margin = Twips(280)
        section.left_margin = Twips(1700)

    def _setup_styles(self) -> None:
        style = self.document.styles["Normal"]
        style.font.name = BODY_FONT
        style.font.size = BODY_SIZE
        rpr = style.element.get_or_add_rPr().get_or_add_rFonts()
        rpr.set(qn("w:eastAsia"), BODY_FONT)
        rpr.set(qn("w:cs"), BODY_FONT)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0

    # ------------------------------------------------------------------ #
    # Абзацы
    # ------------------------------------------------------------------ #

    def _paragraph(self, align=None, first_line=None, space_before=None):
        paragraph = self.document.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        if first_line is not None:
            paragraph.paragraph_format.first_line_indent = first_line
        if space_before is not None:
            paragraph.paragraph_format.space_before = space_before
        return paragraph

    def empty(self, count: int = 1) -> None:
        for _ in range(count):
            self._paragraph()

    def body(self, text: str, bold_prefix: str | None = None) -> None:
        """Абзац основного текста: по ширине, с абзацным отступом."""
        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.JUSTIFY, FIRST_LINE_INDENT)
        if bold_prefix:
            run = paragraph.add_run(bold_prefix)
            run.bold = True
        paragraph.add_run(text)

    def centered(self, text: str, bold: bool = False) -> None:
        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER)
        run = paragraph.add_run(text)
        run.bold = bold

    def numbered(self, items: list[str]) -> None:
        for index, text in enumerate(items, start=1):
            paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.JUSTIFY, FIRST_LINE_INDENT)
            paragraph.add_run(f"{index}. {text}")

    def page_break(self) -> None:
        self.document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ #
    # Титульный лист
    # ------------------------------------------------------------------ #

    def _title_page(self, number: int, topic_lines: list[str]) -> None:
        self.centered("МИНИСТЕРСТВО ОБРАЗОВАНИЯ РЕСПУБЛИКИ БЕЛАРУСЬ", bold=True)
        self.empty()
        self.centered("УЧРЕЖДЕНИЕ ОБРАЗОВАНИЯ", bold=True)
        self.centered(UNIVERSITY, bold=True)
        self.empty(3)
        self.centered(FACULTY)
        self.empty()
        self.centered(f"ОТЧЁТ ПО ЛАБОРАТОРНОЙ РАБОТЕ № {number}")
        self.centered("по дисциплине:")
        self.centered(f"«{DISCIPLINE}»")
        self.centered("на тему:")
        for line in topic_lines:
            self.centered(line)

        self.empty(12)

        for text in (STUDENT, "Принял:", TEACHER):
            paragraph = self._paragraph()
            paragraph.paragraph_format.left_indent = Cm(10.0)
            paragraph.add_run(text)

        self.empty(7)
        self.centered(CITY_YEAR)
        self.page_break()

    # ------------------------------------------------------------------ #
    # Рисунки
    # ------------------------------------------------------------------ #

    def figure(self, images: list[Path], caption: str, tablet: bool = False) -> None:
        """Вставляет изображение (одно или склеенную пару) с подписью."""
        composed = _compose(images)

        width = TABLET_PAIR_WIDTH if tablet else PHONE_PAIR_WIDTH
        ratio = composed.height / composed.width
        height = Cm(width.cm * ratio)
        if not tablet and height > PHONE_PAIR_MAX_HEIGHT:
            height = PHONE_PAIR_MAX_HEIGHT
            width = Cm(height.cm / ratio)

        buffer = BytesIO()
        composed.save(buffer, format="PNG")
        buffer.seek(0)

        self.empty()
        paragraph = self._paragraph(WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.add_run().add_picture(buffer, width=width, height=height)
        self.centered(caption)
        self.empty()

    # ------------------------------------------------------------------ #
    # Приложение с листингами
    # ------------------------------------------------------------------ #

    def appendix(self, files: list[str], ref: str | None = None) -> None:
        """Листинги программы.

        `ref` — ветка репозитория, из которой берётся исходный код. Это
        позволяет приложить к отчёту состояние проекта именно на момент
        сдачи соответствующей лабораторной работы.
        """
        self.empty()
        self.centered("ПРИЛОЖЕНИЕ А")
        self.centered("(обязательное)")
        self.empty()
        self.centered("Листинг программы")
        self.empty()

        for relative in files:
            paragraph = self._paragraph()
            run = paragraph.add_run(f"{Path(relative).name}:")
            run.bold = True
            run.italic = True

            text = _read_source(relative, ref)
            for line in text.replace("\r\n", "\n").rstrip("\n").split("\n"):
                code = self._paragraph()
                code.paragraph_format.space_after = Pt(0)
                run = code.add_run(line)
                run.font.size = CODE_SIZE
            self.empty()

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(str(path))
        print(f"Готово: {path.relative_to(ROOT)}")


def _read_source(relative: str, ref: str | None) -> str:
    """Читает файл из рабочего каталога или из указанной ветки репозитория."""
    if ref is None:
        return (ROOT / relative).read_text(encoding="utf-8")

    result = subprocess.run(
        ["git", "show", f"{ref}:{relative}"],
        cwd=ROOT, capture_output=True, check=True,
    )
    return result.stdout.decode("utf-8")


def _compose(images: list[Path]) -> Image.Image:
    """Склеивает снимки экрана в одно изображение по горизонтали."""
    opened = [Image.open(path).convert("RGB") for path in images]
    if len(opened) == 1:
        return opened[0]

    height = max(image.height for image in opened)
    scaled = [
        image if image.height == height
        else image.resize((round(image.width * height / image.height), height),
                          Image.LANCZOS)
        for image in opened
    ]

    gap = round(height * 0.02)
    total = sum(image.width for image in scaled) + gap * (len(scaled) - 1)
    canvas = Image.new("RGB", (total, height), "white")

    offset = 0
    for image in scaled:
        canvas.paste(image, (offset, 0))
        offset += image.width + gap
    return canvas


def screens(lab: str, *names: str) -> list[Path]:
    """Пути к снимкам экрана лабораторной работы."""
    return [ROOT / lab / "screens" / f"{name}.png" for name in names]
